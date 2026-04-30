import asyncio
import pandas as pd
from playwright.async_api import async_playwright
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from datetime import datetime

CSV_FILE = "export_sfmc.csv"
OUTPUT_FILE = "resultats_agences.xlsx"
DELAI_ENTRE_REQUETES = 2  # secondes entre chaque test (évite le blocage par Google)


async def tester_lien(page, url):
    if not isinstance(url, str) or url.strip() == "":
        return "URL MANQUANTE"

    try:
        await page.goto(url.strip(), timeout=20000, wait_until="domcontentloaded")
        await page.wait_for_timeout(4000)

        contenu = await page.inner_text("body")

        # Détection en français et en anglais car Google Maps adapte la langue selon la région
        if any(x in contenu for x in ["Fermé définitivement", "Permanently closed", "définitivement fermé"]):
            return "FERMEE DEFINITIVEMENT"
        elif any(x in contenu for x in ["Fermé temporairement", "Temporarily closed"]):
            return "FERMEE TEMPORAIREMENT"
        elif any(x in contenu for x in ["Ouvert maintenant", "Open now", "Ouvert ⋅", "Opens"]):
            return "OUVERTE"
        elif any(x in contenu for x in ["Fermé ⋅", "Closes", "Ferme à"]):
            return "OUVERTE (horaires trouvés)"
        elif "maps.google" in page.url or "google.com/maps" in page.url:
            return "PAGE MAPS (statut non détecté)"
        else:
            return "STATUT INCONNU"

    except Exception as e:
        return f"ERREUR ({str(e)[:50]})"


async def tester_toutes_agences(df):
    resultats = []

    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=True)
        context = await browser.new_context(
            user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
        )
        page = await context.new_page()

        total = len(df)
        for i, row in df.iterrows():
            url = row.get("URL_Avis", "")
            agence = row.get("Agence", f"Ligne {i+1}")
            localite = row.get("Localite", "")
            departement = row.get("Departement", "")

            print(f"[{i+1}/{total}] Test : {agence} ({localite})...")

            statut = await tester_lien(page, url)
            print(f"         -> {statut}")

            resultats.append({
                "Departement": departement,
                "CodePostal": row.get("CodePostal", ""),
                "Localite": localite,
                "Agence": agence,
                "NomEntreprise": row.get("NomEntreprise", ""),
                "URL_Avis": url,
                "Statut": statut,
            })

            await asyncio.sleep(DELAI_ENTRE_REQUETES)

        await browser.close()

    return resultats


def generer_rapport(resultats):
    wb = Workbook()
    ws = wb.active
    ws.title = "Résultats"

    headers = ["Département", "Code Postal", "Localité", "Agence", "Nom Entreprise", "URL Avis", "Statut"]
    header_fill = PatternFill("solid", start_color="1F4E79")
    header_font = Font(bold=True, color="FFFFFF", name="Arial", size=11)

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.row_dimensions[1].height = 25

    couleurs = {
        "OUVERTE":               "C6EFCE",  # vert
        "FERMEE DEFINITIVEMENT": "FFCCCC",  # rouge
        "FERMEE TEMPORAIREMENT": "FFEB9C",  # orange
        "URL MANQUANTE":         "FFEB9C",  # orange
        "ERREUR":                "FFB3B3",  # rouge foncé
    }

    for row_idx, r in enumerate(resultats, 2):
        valeurs = [
            r["Departement"], r["CodePostal"], r["Localite"],
            r["Agence"], r["NomEntreprise"], r["URL_Avis"], r["Statut"]
        ]

        statut = r["Statut"]
        bg = next((v for k, v in couleurs.items() if statut.startswith(k)), "EDEDED")
        fill = PatternFill("solid", start_color=bg)

        for col_idx, val in enumerate(valeurs, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.fill = fill
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="center")

    largeurs = [12, 12, 20, 30, 40, 50, 28]
    for i, w in enumerate(largeurs, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("Résumé")
    total = len(resultats)
    ouvertes = sum(1 for r in resultats if r["Statut"].startswith("OUVERTE"))
    fermees_def = sum(1 for r in resultats if r["Statut"] == "FERMEE DEFINITIVEMENT")
    fermees_tmp = sum(1 for r in resultats if r["Statut"] == "FERMEE TEMPORAIREMENT")
    erreurs = sum(1 for r in resultats if r["Statut"].startswith("ERREUR") or r["Statut"] in ("STATUT INCONNU", "PAGE MAPS (statut non détecté)", "URL MANQUANTE"))

    resume_data = [
        ("Rapport généré le", datetime.now().strftime("%d/%m/%Y %H:%M")),
        ("", ""),
        ("TOTAL agences testées", total),
        ("Ouvertes", ouvertes),
        ("Fermées définitivement", fermees_def),
        ("Fermées temporairement", fermees_tmp),
        ("Statut inconnu / Erreur", erreurs),
    ]

    for row_idx, (label, valeur) in enumerate(resume_data, 1):
        ws2.cell(row=row_idx, column=1, value=label).font = Font(bold=True, name="Arial")
        ws2.cell(row=row_idx, column=2, value=valeur).font = Font(name="Arial")

    ws2.column_dimensions["A"].width = 35
    ws2.column_dimensions["B"].width = 25

    wb.save(OUTPUT_FILE)
    print(f"\nRapport Excel généré : {OUTPUT_FILE}")


async def main():
    print("=" * 60)
    print("  Vérificateur d'agences SFMC — Google Maps")
    print("=" * 60)

    try:
        df = pd.read_csv(CSV_FILE, encoding="utf-8-sig", sep=",")
    except FileNotFoundError:
        print(f"\nFichier introuvable : {CSV_FILE}")
        print("   -> Exporte ta DE depuis SFMC et renomme-le 'export_sfmc.csv'")
        return
    except Exception as e:
        print(f"\nErreur de lecture : {e}")
        return

    print(f"\n{len(df)} agences chargées depuis {CSV_FILE}")
    print("Démarrage des tests...\n")

    resultats = await tester_toutes_agences(df)
    generer_rapport(resultats)

    ouvertes = sum(1 for r in resultats if r["Statut"].startswith("OUVERTE"))
    fermees = sum(1 for r in resultats if r["Statut"] == "FERMEE DEFINITIVEMENT")
    print(f"\nRESUME : {ouvertes} ouvertes | {fermees} fermées définitivement | {len(resultats) - ouvertes - fermees} autres")


if __name__ == "__main__":
    asyncio.run(main())
